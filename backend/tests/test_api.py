import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from app.services.data_generator import PRESETS, generate_dataset, write_dataset

SESSION = {
    "name": "End Semester Examination", "academic_year": "2026-27", "semester": "VII",
    "start_date": "2026-11-20", "end_date": "2026-12-31", "slots_per_day": 2,
}
FAST_GA = {"population_size": 20, "generations": 8, "early_stopping": False}


@pytest.fixture
def client(tmp_path):
    app = create_app(f"sqlite:///{(tmp_path / 'test.db').as_posix()}", job_mode="inline")
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def csv_dir(tmp_path_factory):
    path = tmp_path_factory.mktemp("csv")
    write_dataset(generate_dataset(PRESETS["small"]), path)
    return path


@pytest.fixture
def sid(client):
    return client.post("/api/sessions", json=SESSION).json()["id"]


@pytest.fixture
def loaded(client, sid):
    r = client.post(f"/api/sessions/{sid}/load-sample", json={"name": "small"})
    assert r.status_code == 200
    return sid


def upload(client, url, csv_dir, name, text=None):
    data = text.encode() if text is not None else (csv_dir / name).read_bytes()
    return client.post(url, files={"file": (name, data, "text/csv")})


def run_ga(client, sid, **params):
    r = client.post("/api/optimization/run", json={"session_id": sid, "params": {**FAST_GA, **params}})
    assert r.status_code == 202, r.text
    return r.json()


# --- sessions and data --------------------------------------------------------------------------
def test_health_and_session_crud(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    r = client.post("/api/sessions", json=SESSION)
    assert r.status_code == 201 and r.json()["slots_per_day"] == 2
    sid = r.json()["id"]
    assert client.get(f"/api/sessions/{sid}").json()["counts"]["exams"] == 0
    assert [s["id"] for s in client.get("/api/sessions").json()] == [sid]
    assert client.get("/api/sessions/999").status_code == 404


def test_session_validation(client):
    bad = {**SESSION, "end_date": "2026-11-01"}
    assert client.post("/api/sessions", json=bad).status_code == 422
    assert client.post("/api/sessions", json={**SESSION, "slots_per_day": 9}).status_code == 422


def test_load_sample_and_validate_data(client, loaded):
    detail = client.get(f"/api/sessions/{loaded}").json()
    assert detail["counts"]["exams"] == 20 and detail["counts"]["slots"] == 14
    check = client.post(f"/api/sessions/{loaded}/validate").json()
    assert check["ok"] and not check["errors"]
    assert check["stats"]["exams"] == 20 and check["stats"]["greedy_slots_needed"] <= check["stats"]["slots"]


def test_validate_reports_missing_data(client, sid):
    check = client.post(f"/api/sessions/{sid}/validate").json()
    assert not check["ok"] and "no exams imported yet" in check["errors"]


def test_csv_import_flow_end_to_end(client, sid, csv_dir):
    base = f"/api/sessions/{sid}"
    assert upload(client, f"{base}/exams/import", csv_dir, "exams.csv").json()["imported"] == 20
    assert upload(client, f"{base}/students/import", csv_dir, "students.csv").json()["imported"] == 200
    assert upload(client, f"{base}/rooms/import", csv_dir, "rooms.csv").json()["imported"] == 8
    enrol = upload(client, f"{base}/enrollments/import", csv_dir, "enrollments.csv").json()
    assert enrol["imported"] > 0
    gen = client.post(f"{base}/slots/generate").json()
    assert gen["imported"] == 60  # 30 weekdays from 20 Nov to 31 Dec 2026, 2 slots/day
    assert len(client.get(f"{base}/exams").json()) == 20
    assert client.post(f"{base}/validate").json()["ok"]


def test_import_errors_are_422_with_messages(client, sid, csv_dir):
    base = f"/api/sessions/{sid}"
    dup = "exam_id,subject_code,subject_name,department,semester,duration_minutes,student_count,exam_type\n" \
          "E1,C,N,D,V,180,10,theory\nE1,C,N,D,V,180,10,theory\n"
    r = upload(client, f"{base}/exams/import", csv_dir, "exams.csv", dup)
    assert r.status_code == 422 and any("duplicate exam_id 'E1'" in e for e in r.json()["detail"]["errors"])
    # enrollments before exams/students: references cannot resolve
    r = upload(client, f"{base}/enrollments/import", csv_dir, "enrollments.csv")
    assert r.status_code == 422 and "does not exist" in r.json()["detail"]["errors"][0]
    r = upload(client, f"{base}/rooms/import", csv_dir, "rooms.csv", "room_id,room_code,capacity\nR1,A1,0\n")
    assert r.status_code == 422 and "room R1 has non-positive capacity (0)." in r.json()["detail"]["errors"][0]
    r = upload(client, f"{base}/exams/import", csv_dir, "exams.csv", "")
    assert r.status_code == 422


def test_changing_data_invalidates_timetables(client, loaded, csv_dir):
    run = run_ga(client, loaded)
    tid = run["timetable_id"]
    r = upload(client, f"/api/sessions/{loaded}/rooms/import", csv_dir, "rooms.csv")
    assert r.json()["invalidated_timetables"] == 1
    assert client.get(f"/api/timetables/{tid}").status_code == 404
    assert client.get(f"/api/optimization/{run['id']}").json()["timetable_id"] is None  # run history kept


# --- optimization ----------------------------------------------------------------------------------
def test_ga_run_produces_a_valid_timetable_with_progress_history(client, loaded):
    run = run_ga(client, loaded)
    assert run["status"] == "completed" and run["hard_violations"] == 0
    assert run["initial_penalty"] > run["best_penalty"]
    assert run["generations_run"] == 8 and run["dataset_hash"] and run["seed"] == 42
    progress = client.get(f"/api/optimization/{run['id']}/progress").json()
    assert progress["status"] == "completed" and len(progress["history"]) == 9
    assert progress["history"][-1]["best_penalty"] == pytest.approx(run["best_penalty"])
    assert [r["id"] for r in client.get(f"/api/sessions/{loaded}/runs").json()] == [run["id"]]


def test_same_seed_reproduces_through_the_api(client, loaded):
    a, b = run_ga(client, loaded), run_ga(client, loaded)
    assert a["best_penalty"] == b["best_penalty"]
    c = run_ga(client, loaded, seed=7)
    assert c["seed"] == 7


def test_run_is_blocked_without_data_or_when_infeasible(client, sid, csv_dir):
    r = client.post("/api/optimization/run", json={"session_id": sid, "params": FAST_GA})
    assert r.status_code == 422 and "no exams imported yet" in r.json()["detail"]["errors"]
    client.post(f"/api/sessions/{sid}/load-sample", json={"name": "small"})
    tiny_room = "room_id,room_code,capacity\nR1,A1,5\nR2,A2,6\n"
    upload(client, f"/api/sessions/{sid}/rooms/import", csv_dir, "rooms.csv", tiny_room)
    r = client.post("/api/optimization/run", json={"session_id": sid, "params": FAST_GA})
    assert r.status_code == 422 and any("seats but the largest" in e for e in r.json()["detail"]["errors"])


def test_invalid_ga_parameters_rejected(client, loaded):
    r = client.post("/api/optimization/run", json={"session_id": loaded, "params": {"population_size": 10, "elite_count": 10}})
    assert r.status_code == 422
    r = client.post("/api/optimization/run", json={"session_id": loaded, "method": "nonsense"})
    assert r.status_code == 422
    assert client.get("/api/optimization/999").status_code == 404


def test_baseline_endpoint(client, loaded):
    r = client.post("/api/baseline/run", json={"session_id": loaded, "method": "greedy_cost_aware"})
    assert r.status_code == 202 and r.json()["status"] == "completed" and r.json()["hard_violations"] == 0
    ga = run_ga(client, loaded, generations=30)
    assert ga["best_penalty"] < r.json()["best_penalty"]  # the GA beats the strong greedy baseline


# --- timetables -------------------------------------------------------------------------------------
def test_timetable_detail_filters_and_validation(client, loaded):
    tid = run_ga(client, loaded)["timetable_id"]
    detail = client.get(f"/api/timetables/{tid}").json()
    assert detail["status"] == "draft" and detail["total_entries"] == 20 and len(detail["entries"]) == 20
    assert detail["fitness"] == pytest.approx(1 / (1 + detail["penalty"]))
    dept = detail["entries"][0]["department"]
    by_dept = client.get(f"/api/timetables/{tid}", params={"department": dept}).json()
    assert 0 < len(by_dept["entries"]) <= 20 and all(e["department"] == dept for e in by_dept["entries"])
    day = detail["entries"][0]["date"]
    by_day = client.get(f"/api/timetables/{tid}", params={"date": day}).json()
    assert all(e["date"] == day for e in by_day["entries"])
    code = detail["entries"][0]["subject_code"]
    assert len(client.get(f"/api/timetables/{tid}", params={"subject": code.lower()}).json()["entries"]) >= 1
    v = client.post(f"/api/timetables/{tid}/validate").json()
    assert v["valid"] and v["hard_violations"] == 0 and v["student_conflicts"] == 0


def test_manually_edited_invalid_timetable_is_rejected(client, loaded):
    tid = run_ga(client, loaded)["timetable_id"]
    entries = client.get(f"/api/timetables/{tid}").json()["entries"]
    edited = [{"exam_id": e["exam_id"], "slot_id": entries[0]["slot_id"], "room_id": e["room_id"]} for e in entries]
    r = client.post("/api/timetables/validate", json={"session_id": loaded, "entries": edited})
    assert r.status_code == 200 and not r.json()["valid"] and r.json()["student_conflicts"] > 0
    partial = edited[:3]
    r = client.post("/api/timetables/validate", json={"session_id": loaded, "entries": partial}).json()
    assert r["unassigned"] == 17


def test_approval_supersedes_previous_and_rejects_infeasible(client, loaded):
    first = run_ga(client, loaded)["timetable_id"]
    second = run_ga(client, loaded, seed=9)["timetable_id"]
    assert client.post(f"/api/timetables/{first}/approve").json()["status"] == "approved"
    assert client.post(f"/api/timetables/{second}/approve").json()["status"] == "approved"
    statuses = {t["id"]: t["status"] for t in client.get(f"/api/sessions/{loaded}/timetables").json()}
    assert statuses == {first: "superseded", second: "approved"}

    bad = client.post("/api/baseline/run", json={"session_id": loaded, "method": "random_feasible"}).json()
    assert bad["hard_violations"] > 0
    bad_tt = client.get(f"/api/timetables/{bad['timetable_id']}").json()
    assert bad_tt["status"] == "infeasible"
    assert client.post(f"/api/timetables/{bad['timetable_id']}/approve").status_code == 409
    assert client.get(f"/api/export/{bad['timetable_id']}/csv").status_code == 409


def test_exports(client, loaded):
    tid = run_ga(client, loaded)["timetable_id"]
    csv_resp = client.get(f"/api/export/{tid}/csv")
    assert csv_resp.status_code == 200 and "attachment" in csv_resp.headers["content-disposition"]
    lines = csv_resp.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "Date,Day,Start,End,Subject Code,Subject,Department,Semester,Exam Type,Room,Building,Students"
    assert len(lines) == 21
    xlsx = client.get(f"/api/export/{tid}/xlsx")
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws.max_row == 21 and ws["E1"].value == "Subject Code"
    assert client.get("/api/export/999/csv").status_code == 404


# --- experiments ---------------------------------------------------------------------------------
def test_experiments_listing_and_errors(client):
    listing = client.get("/api/experiments").json()
    assert {"main", "sensitivity", "ablation", "scaling"} <= {e["name"] for e in listing}
    assert client.get("/api/experiments/not_a_real_one").status_code == 404
    assert client.post("/api/experiments/run", json={"name": "../etc/passwd"}).status_code == 404
    assert client.get("/api/experiments/jobs/missing").status_code == 404
