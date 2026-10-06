"""Timetable persistence, presentation, validation, approval and export."""
from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select, update
from sqlalchemy.orm import Session as DbSession

from app import models as orm
from app.core.domain import Assignment
from app.optimization.chromosome import Chromosome
from app.optimization.context import OptimizationContext
from app.optimization.fitness import Evaluation
from app.optimization.validator import ValidationResult, validate_timetable
from app.services.dataset_service import LoadedDataset, load_dataset_from_db

EXPORT_COLUMNS = [
    "Date", "Day", "Start", "End", "Subject Code", "Subject", "Department",
    "Semester", "Exam Type", "Room", "Building", "Students",
]


def persist_chromosome(
    db: DbSession,
    session_id: int,
    run_id: int | None,
    name: str,
    loaded: LoadedDataset,
    ctx: OptimizationContext,
    chromosome: Chromosome,
    evaluation: Evaluation,
) -> orm.Timetable:
    timetable = orm.Timetable(
        session_id=session_id,
        run_id=run_id,
        name=name,
        fitness=evaluation.fitness,
        penalty=evaluation.penalty,
        hard_violations=evaluation.hard_violations,
        status="draft" if evaluation.hard_violations == 0 else "infeasible",
        breakdown=evaluation.as_dict(),
    )
    db.add(timetable)
    db.flush()
    db.add_all(
        orm.TimetableEntry(
            timetable_id=timetable.id,
            exam_id=loaded.exam_pk[a.exam_id],
            slot_id=loaded.slot_pk[a.slot_id],
            room_id=loaded.room_pk[a.room_id],
        )
        for a in chromosome.assignments(ctx)
    )
    db.flush()
    return timetable


def entry_rows(db: DbSession, timetable_id: int) -> list[dict]:
    """Entries joined with exam/slot/room details, ordered by date, time and room."""
    stmt = (
        select(orm.TimetableEntry, orm.Exam, orm.Slot, orm.Room)
        .join(orm.Exam, orm.Exam.id == orm.TimetableEntry.exam_id)
        .join(orm.Slot, orm.Slot.id == orm.TimetableEntry.slot_id)
        .join(orm.Room, orm.Room.id == orm.TimetableEntry.room_id)
        .where(orm.TimetableEntry.timetable_id == timetable_id)
        .order_by(orm.Slot.date, orm.Slot.start_time, orm.Room.room_code)
    )
    return [
        {
            "exam_id": e.external_id, "subject_code": e.subject_code, "subject_name": e.subject_name,
            "department": e.department, "semester": e.semester, "exam_type": e.exam_type,
            "students": e.student_count, "duration_minutes": e.duration_minutes,
            "slot_id": s.external_id, "date": s.date, "start_time": s.start_time, "end_time": s.end_time,
            "room_id": r.external_id, "room_code": r.room_code, "building": r.building, "capacity": r.capacity,
        }
        for _, e, s, r in db.execute(stmt).all()
    ]


def filter_rows(
    rows: list[dict],
    department: str | None = None,
    semester: str | None = None,
    on_date: date | None = None,
    room: str | None = None,
    subject: str | None = None,
) -> list[dict]:
    """PRD section 10 filters; subject matches code or name, case-insensitively."""
    out = rows
    if department:
        out = [r for r in out if r["department"].lower() == department.lower()]
    if semester:
        out = [r for r in out if r["semester"].lower() == semester.lower()]
    if on_date:
        out = [r for r in out if r["date"] == on_date]
    if room:
        out = [r for r in out if room.lower() in (r["room_code"].lower(), r["room_id"].lower())]
    if subject:
        needle = subject.lower()
        out = [r for r in out if needle in r["subject_code"].lower() or needle in r["subject_name"].lower()]
    return out


def validate_stored(db: DbSession, timetable: orm.Timetable) -> ValidationResult:
    """Run the independent validator over a stored timetable using the session's current data."""
    loaded = load_dataset_from_db(db, timetable.session_id)
    rows = entry_rows(db, timetable.id)
    return validate_timetable(loaded.dataset, [Assignment(r["exam_id"], r["slot_id"], r["room_id"]) for r in rows])


def approve(db: DbSession, timetable: orm.Timetable) -> ValidationResult:
    """Approve a timetable only if the validator finds no hard violations; supersedes the previous one."""
    result = validate_stored(db, timetable)
    if not result.valid:
        return result
    db.execute(
        update(orm.Timetable)
        .where(orm.Timetable.session_id == timetable.session_id, orm.Timetable.status == "approved")
        .values(status="superseded")
    )
    timetable.status = "approved"
    db.flush()
    return result


def _export_table(rows: list[dict]) -> list[list]:
    return [
        [
            r["date"].isoformat(), r["date"].strftime("%a"), r["start_time"].strftime("%H:%M"),
            r["end_time"].strftime("%H:%M"), r["subject_code"], r["subject_name"], r["department"],
            r["semester"], r["exam_type"], r["room_code"], r["building"], r["students"],
        ]
        for r in rows
    ]


def export_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXPORT_COLUMNS)
    writer.writerows(_export_table(rows))
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8


def export_xlsx(rows: list[dict], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.append(EXPORT_COLUMNS)
    for line in _export_table(rows):
        ws.append(line)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")
    for idx, header in enumerate(EXPORT_COLUMNS, 1):
        width = max([len(header)] + [len(str(row[idx - 1])) for row in _export_table(rows)])
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 48)
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "0072B2"
    wb.properties.title = title
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
